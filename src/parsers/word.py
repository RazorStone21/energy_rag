"""读取 DOCX 正文和表格，保留标题路径及原文顺序，不推算排版页码。

本阶段处理文档主体；不提取图片描述、页眉页脚、脚注、文本框和修订内容。
表格转为 Markdown，合并单元格按网格重复内容，嵌套表格转为所在单元格内的文本。
"""

from __future__ import annotations

import re
from html import escape
from pathlib import Path

from ..schemas import ParseResult

# 章节标题的长度上限；超过它的「标题」实际上是被误套了标题样式的正文。
_MAX_HEADING_CHARS = 40


def _looks_like_heading(text):
    """判断被套了标题样式的段落是否真的像标题。

    真标题都很短，也不含句末标点；被误套标题样式的正文则是成句的，
    必然带「。」。只按长度筛不够——实测有 30 字的正文片段躲过了长度上限。
    """
    text = text.strip()
    return len(text) <= _MAX_HEADING_CHARS and "。" not in text


def _outline_level(element):
    """读取段落或样式显式设置的大纲级别；9 表示正文而非标题。"""
    levels = element.xpath("./w:pPr/w:outlineLvl/@w:val")
    if not levels:
        return None
    return int(levels[0])


def _heading_level(paragraph):
    """识别大纲级别、内置标题及其继承样式，不把普通加粗文字猜成标题。"""
    level = _outline_level(paragraph._p)
    if level is not None:
        return level + 1 if 0 <= level <= 8 else None
    style = paragraph.style
    visited = set()
    while style is not None and style.style_id not in visited:
        visited.add(style.style_id)
        level = _outline_level(style.element)
        if level is not None:
            return level + 1 if 0 <= level <= 8 else None
        name = style.name or ""
        if name.lower() == "title" or name == "标题":
            return 0
        match = re.fullmatch(r"(?:heading|标题)\s*([1-9])", name, flags=re.IGNORECASE)
        if match:
            return int(match.group(1))
        style = style.base_style
    return None


def _cell_text(cell):
    """按单元格内部顺序读取段落和嵌套表格，避免遗漏嵌套表格中的文字。"""
    from docx.table import Table

    parts = []
    for block in cell.iter_inner_content():
        if isinstance(block, Table):
            content = _table_markdown(block)
        else:
            content = block.text
        if content.strip():
            parts.append(content)
    return "\n".join(parts)


def _table_markdown(table):
    """按布局网格输出全部数据行，用通用列名作表头，不假定第一行一定是表头。"""
    rows = []
    for row in table.rows:
        cells = [""] * row.grid_cols_before
        for cell in row.cells:
            content = _cell_text(cell).strip()
            # 防止原文中的竖线和换行打断 Markdown 的列与行。
            content = escape(content).replace("\\", "\\\\").replace("|", "\\|")
            cells.append(content.replace("\n", "<br>"))
        cells.extend([""] * row.grid_cols_after)
        rows.append(cells)
    if not rows or not any(any(row) for row in rows):
        return ""
    width = max(len(row) for row in rows)
    header = [f"第{index}列" for index in range(1, width + 1)]
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    for row in rows:
        padded = row + [""] * (width - len(row))
        lines.append("| " + " | ".join(padded) + " |")
    return "\n".join(lines)


class WordParser:
    """解析 DOCX，正文按章节和表格边界分组，表格保持完整。"""

    def __init__(self, document_factory=None):
        """保存文档创建函数；python-docx 和 LangChain 在实际解析时才导入。"""
        self.document_factory = document_factory

    def _make_document(self, content, path, headings, **metadata):
        """创建检索文档，统一保存文件名、标题路径和当前元素的位置信息。"""
        factory = self.document_factory
        if factory is None:
            from langchain_core.documents import Document

            factory = Document
        return factory(
            page_content=content,
            metadata={
                "source": path.name,
                "heading_path": " > ".join(title for _, title in headings),
                **metadata,
            },
        )

    def _append_text(self, result, paragraphs, path, headings):
        """将连续段落放入一个正文块，记录其原文段落范围，随后清空待处理段落。"""
        if not paragraphs:
            return
        result.texts.append(
            self._make_document(
                "\n\n".join(text for _, _, text in paragraphs),
                path,
                headings,
                type="text",
                block_kind="section",
                block_index=paragraphs[0][0],
                paragraph_start=paragraphs[0][1],
                paragraph_end=paragraphs[-1][1],
            )
        )
        paragraphs.clear()

    def _parse_document(self, word, path):
        """依次处理主体段落和表格；标题或表格出现时结束上一组正文。"""
        from docx.table import Table

        result = ParseResult()
        headings, paragraphs = [], []
        paragraph_index = table_index = 0
        for block_index, block in enumerate(word.iter_inner_content(), start=1):
            if isinstance(block, Table):
                self._append_text(result, paragraphs, path, headings)
                table_index += 1
                content = _table_markdown(block)
                if content:
                    result.tables.append(
                        self._make_document(
                            content,
                            path,
                            headings,
                            type="table",
                            block_kind="table",
                            block_index=block_index,
                            table_index=table_index,
                        )
                    )
                continue
            paragraph_index += 1
            text = block.text
            if not text.strip():
                continue
            level = _heading_level(block)
            # 样式可靠但会被误用（实测有文档把整段正文标成 Heading 3），
            # 因此还要看内容是否真的像标题，见 _looks_like_heading。
            if level is not None and _looks_like_heading(text):
                self._append_text(result, paragraphs, path, headings)
                while headings and headings[-1][0] >= level:
                    headings.pop()
                headings.append((level, text.strip()))
            paragraphs.append((block_index, paragraph_index, text))
        self._append_text(result, paragraphs, path, headings)
        if not result.texts and not result.tables:
            result.errors.append("word: 文档主体中没有可用正文或表格")
        return result

    def parse(self, path: str | Path, on_status=None) -> ParseResult:
        """读取 DOCX；格式不支持、文件损坏或解析失败时返回错误，供入库流程保留旧数据。

        on_status 是解析器协议要求的进度回调；DOCX 没有需要分步上报的环节，
        这里接受但不使用，这样注册表可以用同一种方式调用所有解析器。
        """
        path = Path(path)
        if path.suffix.lower() != ".docx":
            return ParseResult(errors=["word: 当前仅支持 .docx，请先将旧版 .doc 转换为 .docx"])
        try:
            from docx import Document

            return self._parse_document(Document(str(path)), path)
        except Exception as exc:
            return ParseResult(errors=[f"word: {exc}"])
