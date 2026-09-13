"""读取 XLSX 工作表，按连续数据区域分片，保留表头、公式及原文行列位置。

合并单元格使用左上角的值展开；包含隐藏工作表和隐藏行列，不执行公式或读取图表图片。
读取公式和缓存各打开一次工作簿，始终不保存或修改原文件。
"""

from __future__ import annotations

from contextlib import ExitStack
from datetime import date, datetime, time
from html import escape
from pathlib import Path

from ..config import ExcelSettings
from ..schemas import ParseResult


def _format_value(value):
    """将原始单元格值转为文字，保留零、布尔值和日期，不套用 Excel 显示格式。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value).strip()


def _cell_text(cell, cached_cell):
    """普通单元格读取原值；公式保留表达式及缓存，缺少缓存时明确标注。"""
    if cell.data_type != "f":
        return _format_value(cell.value)
    formula = cell.value
    if not isinstance(formula, str):
        formula = getattr(formula, "text", None) or "表达式不可用"
    if cached_cell.value is None:
        return f"公式：{formula}（无缓存结果）"
    return f"{_format_value(cached_cell.value)}（缓存结果；公式：{formula}）"


def _escape_cell(value):
    """转义竖线、反斜线和换行，防止单元格内容改变 Markdown 表格结构。"""
    return (
        escape(value)
        .replace("\\", "\\\\")
        .replace("|", "\\|")
        .replace("\r\n", "\n")
        .replace("\n", "<br>")
    )


class ExcelParser:
    """将 XLSX 的各工作表转为表格片段，每片记录工作表名、行号和列范围。"""

    def __init__(self, settings: ExcelSettings | None = None, document_factory=None):
        """保存分片配置与文档创建函数，openpyxl 在真正解析时才导入。"""
        self.settings = settings or ExcelSettings()
        self.document_factory = document_factory

    def _sheet_regions(self, sheet, cached_sheet):
        """展开合并单元格，按全空行分隔数据区域；隐藏行列也参与读取。"""
        if sheet.max_row * sheet.max_column > self.settings.max_sheet_cells:
            raise ValueError(
                f"工作表 {sheet.title} 超出 excel.max_sheet_cells，请清理多余格式或调整上限"
            )
        merged_sources = {}
        for area in sheet.merged_cells.ranges:
            for row in range(area.min_row, area.max_row + 1):
                for column in range(area.min_col, area.max_col + 1):
                    merged_sources[row, column] = (area.min_row, area.min_col)

        region = []
        for row_index, cells in enumerate(sheet.iter_rows(), start=1):
            values = []
            for cell in cells:
                source_row, source_column = merged_sources.get(
                    (cell.row, cell.column), (cell.row, cell.column)
                )
                value = _cell_text(
                    sheet.cell(source_row, source_column),
                    cached_sheet.cell(source_row, source_column),
                )
                values.append(value)
            if any(values):
                region.append((row_index, values))
            elif region:
                yield region
                region = []
        if region:
            yield region

    def _region_documents(self, region, path, sheet):
        """裁去区域两侧空列，按配置合并多行表头并分片，所有数据行保留原始行号。"""
        from openpyxl.utils import get_column_letter

        used_columns = [
            index
            for index in range(len(region[0][1]))
            if any(values[index] for _, values in region)
        ]
        left, right = min(used_columns), max(used_columns) + 1
        header_count = min(self.settings.header_rows, len(region))
        header_rows = region[:header_count]
        data_rows = region[header_count:]
        headers = []
        for column in range(left, right):
            labels = []
            for _, values in header_rows:
                if values[column] and (not labels or labels[-1] != values[column]):
                    labels.append(values[column])
            headers.append(" / ".join(labels) or f"{get_column_letter(column + 1)}列")

        # 只有表头的区域仍保留，避免单行工作表被当成空文件。
        if data_rows:
            starts = range(0, len(data_rows), self.settings.rows_per_chunk)
        else:
            starts = [0]
        for start in starts:
            rows = data_rows[start : start + self.settings.rows_per_chunk]
            lines = [
                f"工作表：{sheet.title}",
                "| 原文行 | " + " | ".join(_escape_cell(header) for header in headers) + " |",
                "| --- | " + " | ".join(["---"] * len(headers)) + " |",
            ]
            for row_index, values in rows:
                lines.append(
                    f"| {row_index} | "
                    + " | ".join(_escape_cell(value) for value in values[left:right])
                    + " |"
                )
            location_rows = rows or header_rows
            metadata = {
                "source": path.name,
                "type": "table",
                "sheet_name": sheet.title,
                "sheet_state": sheet.sheet_state,
                "row_start": location_rows[0][0],
                "row_end": location_rows[-1][0],
                "column_start": get_column_letter(left + 1),
                "column_end": get_column_letter(right),
            }
            if header_rows:
                metadata.update(
                    header_row_start=header_rows[0][0], header_row_end=header_rows[-1][0]
                )
            factory = self.document_factory
            if factory is None:
                from langchain_core.documents import Document

                factory = Document
            yield factory(page_content="\n".join(lines), metadata=metadata)

    def parse(self, path: str | Path, on_status=None) -> ParseResult:
        """读取工作簿；任一工作表失败时记录错误，入库流程不会用部分结果替换旧文件。

        on_status 是解析器协议要求的进度回调；工作表读取没有需要分步上报的环节，
        这里接受但不使用，这样注册表可以用同一种方式调用所有解析器。
        """
        path = Path(path)
        if path.suffix.lower() != ".xlsx":
            return ParseResult(errors=["excel: 当前仅支持 .xlsx，请将旧 .xls 转换为 .xlsx"])
        result = ParseResult()
        try:
            from openpyxl import load_workbook

            # 即使第二次打开或解析中途失败，也关闭已打开的工作簿。
            with ExitStack() as resources:
                formulas = load_workbook(path, data_only=False, keep_links=False)
                resources.callback(formulas.close)
                cached = load_workbook(path, data_only=True, keep_links=False)
                resources.callback(cached.close)
                for sheet in formulas.worksheets:
                    for region in self._sheet_regions(sheet, cached[sheet.title]):
                        result.tables.extend(self._region_documents(region, path, sheet))
        except Exception as exc:
            result.errors.append(f"excel: {exc}")
        if not result.tables and not result.errors:
            result.errors.append("excel: 工作簿中没有可用数据")
        return result
