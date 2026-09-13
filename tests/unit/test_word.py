"""使用临时 DOCX 验证真实段落和表格提取，模型与存储由替身隔离。"""

from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from scripts.export_chunks import _export_record, _location_label, export
from src.chunker import Chunker
from src.config import load_settings
from src.context_builder import ContextBuilder
from src.ingestion import list_documents
from src.parsers.word import WordParser
from src.retrieval.hybrid import rrf_fuse
from src.schemas import SearchHit


def parse_word(tmp_path, word):
    """保存真实 DOCX 再读取，确保验证包含文件格式和 XML 解析过程。"""
    path = tmp_path / "能源报告.DOCX"
    word.save(path)
    parsed = WordParser(document_factory=SimpleNamespace).parse(path)
    assert not parsed.errors, parsed.errors
    return parsed


def test_word_preserves_headings_tables_and_document_order(tmp_path):
    """段落与表格交错时保留顺序，表格后的文字继承标题，新章节结束旧章节。"""
    word = Document()
    word.add_heading("能源报告", level=1)
    word.add_paragraph("前言说明")
    word.add_heading("电力", level=2)
    word.add_paragraph("表格之前")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "发电量"
    table.cell(1, 1).text = "10"
    word.add_paragraph("表格之后")
    word.add_heading("燃气", level=2)
    word.add_paragraph("燃气说明", style="List Bullet")
    parsed = parse_word(tmp_path, word)
    assert len(parsed.texts) == 4 and len(parsed.tables) == 1
    ordered = sorted(parsed.texts + parsed.tables, key=lambda d: d.metadata["block_index"])
    assert [d.metadata["block_index"] for d in ordered] == [1, 3, 5, 6, 7]
    assert ordered[2].metadata["heading_path"] == "能源报告 > 电力"
    assert ordered[3].page_content == "表格之后"
    assert ordered[3].metadata["paragraph_start"] == 5
    assert ordered[-1].metadata["heading_path"] == "能源报告 > 燃气"
    assert "燃气说明" in ordered[-1].page_content
    assert "| 项目 | 数值 |" in parsed.tables[0].page_content
    assert "| 发电量 | 10 |" in parsed.tables[0].page_content
    assert parsed.tables[0].metadata["table_index"] == 1
    assert all("page" not in d.metadata for d in ordered)
    assert not parsed.figures


def test_word_heading_style_inheritance_and_outline_override(tmp_path):
    """自定义样式继承标题级别，段落显式正文大纲可覆盖标题样式，普通加粗不作标题。"""
    word = Document()
    style = word.styles.add_style("自定义能源标题", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = word.styles["Heading 2"]
    word.add_heading("能源", level=1)
    word.add_paragraph("统计", style=style)
    normal = word.add_paragraph("仍是正文", style="Heading 1")
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "9")
    normal._p.get_or_add_pPr().append(outline)
    word.add_paragraph().add_run("普通加粗文字").bold = True
    parsed = parse_word(tmp_path, word)
    assert len(parsed.texts) == 2
    assert parsed.texts[-1].metadata["heading_path"] == "能源 > 统计"
    assert "仍是正文" in parsed.texts[-1].page_content
    assert "普通加粗文字" in parsed.texts[-1].page_content


def test_word_merged_nested_and_escaped_cells_keep_content(tmp_path):
    """横纵合并、嵌套表格、竖线和换行处理后仍保留所有文字。"""
    word = Document()
    table = word.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "合并表头"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "跨行区域"
    table.cell(0, 2).text = "A|B\n第二行"
    cell = table.cell(1, 1)
    cell.text = "嵌套之前"
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "嵌套数据"
    cell.add_paragraph("嵌套之后")
    parsed = parse_word(tmp_path, word)
    assert len(parsed.tables) == 1
    content = parsed.tables[0].page_content
    assert "| 合并表头 | 合并表头 |" in content
    assert content.count("跨行区域") == 2
    assert "A\\|B<br>第二行" in content
    assert content.index("嵌套之前") < content.index("嵌套数据") < content.index("嵌套之后")


def test_word_omitted_grid_cell_does_not_shift_values(tmp_path):
    """行首省略单元格时补齐空位置，后续数值不被错误移到第一列。"""
    word = Document()
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 1).text = "第二列数值"
    row = table.rows[0]
    row._tr.remove(row.cells[0]._tc)
    before = OxmlElement("w:gridBefore")
    before.set(qn("w:val"), "1")
    row._tr.get_or_add_trPr().append(before)
    parsed = parse_word(tmp_path, word)
    assert "|  | 第二列数值 |" in parsed.tables[0].page_content


@pytest.mark.parametrize("kind", ["missing", "corrupt", "empty", "legacy"])
def test_word_invalid_or_empty_files_report_errors(tmp_path, kind):
    """缺失、损坏、无可用正文和旧 DOC 都返回明确错误，不生成可替换的片段。"""
    path = tmp_path / "report.docx"
    if kind == "corrupt":
        path.write_bytes(b"not a docx")
    elif kind == "empty":
        Document().save(path)
    elif kind == "legacy":
        path = path.with_suffix(".doc")
    result = WordParser(document_factory=SimpleNamespace).parse(path)
    assert result.errors and not result.texts and not result.tables


def test_word_discovery_ignores_lock_files_and_legacy_doc(tmp_path):
    """发现大小写 DOCX 后缀，跳过 Word 锁定文件和尚不支持的旧 DOC。"""
    for name in ("report.DOCX", "~$report.docx", "legacy.doc"):
        (tmp_path / name).write_bytes(b"content")
    assert [path.name for path in list_documents(tmp_path, {".docx"})] == ["report.DOCX"]


def test_word_duplicate_text_at_different_positions_is_not_merged(tmp_path):
    """同名章节中相同的正文凭位置区分，跨路命中的同一个片段仍正常合并。"""
    word = Document()
    for _ in range(2):
        word.add_heading("概况", level=1)
        word.add_paragraph("相同正文")
    parsed = parse_word(tmp_path, word)
    assert parsed.texts[0].page_content == parsed.texts[1].page_content
    hits = [SearchHit(d) for d in parsed.texts]
    assert len(rrf_fuse(hits, hits, 10, 60)) == 2


def test_word_chunking_citations_and_exports_keep_locations(tmp_path):
    """正文切分继承段落范围，表格完整保留，提示词与导出都能定位来源。"""
    import csv
    import json

    word = Document()
    word.add_heading("能源报告", level=1)
    word.add_paragraph("用于验证语义切分与来源展示的政策正文。")
    word.add_table(rows=1, cols=1).cell(0, 0).text = "表格数据"
    parsed = parse_word(tmp_path, word)
    settings = load_settings(Path(__file__).resolve().parents[2] / "config.toml")
    splitter = Mock()
    splitter.split_documents.side_effect = lambda docs: docs
    chunks = Chunker(Mock(), settings.splitting, splitter_factory=lambda *a, **k: splitter).split(
        parsed
    )
    assert parsed.tables[0] in chunks
    splitter.split_documents.assert_called_once_with(parsed.texts)
    context = ContextBuilder("{context}\n{question}").build(
        "内容？", [SearchHit(d) for d in chunks]
    )
    assert "标题: 能源报告" in context.prompt
    assert "所属原文段落: 第1—2段" in context.prompt
    assert "第1个表格" in context.prompt
    records = [_export_record(d) for d in chunks]
    assert "第1个表格" in _location_label(records[-1])
    export(records, "json", tmp_path / "chunks.json")
    assert (
        json.loads((tmp_path / "chunks.json").read_text(encoding="utf-8"))[0]["paragraph_start"]
        == 1
    )
    export(records, "csv", tmp_path / "chunks.csv")
    with (tmp_path / "chunks.csv").open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert rows[-1]["table_index"] == "1" and rows[0]["paragraph_end"] == "2"


def test_heading_filter_rejects_body_text_styled_as_heading():
    """被误套标题样式的正文不能进章节路径：真标题短，且不含句末标点。

    真实样本来自 45-green-261789299726.docx：文档把整段正文标成了 Heading 3，
    只按样式判断会把这些半句话写进「标题:」里，再随来源说明进入提示词。
    """
    from src.parsers.word import _looks_like_heading

    assert _looks_like_heading("（二）坚持试点先行，探索形成产业科学发展模式")
    assert _looks_like_heading("（一）政策规范篇")
    assert _looks_like_heading("第一章 总体要求")
    # 实测的两条漏网样本：一条 30 字，长度上限拦不住，靠「。」判定。
    assert not _looks_like_heading("过 100 艘。2025 年，全球新增替代燃料船舶订单 499 艘，占全球")
    assert not _looks_like_heading(
        "速大型船用发动机已实现商业化运行。康明斯、丰田等车用甲醇内燃机技术已较为成熟，"
        "德国、瑞典、丹麦已开展甲醇动力重卡"
    )
