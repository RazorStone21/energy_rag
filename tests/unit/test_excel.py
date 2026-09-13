"""用真实 XLSX 验证工作表、合并单元格、公式缓存和分片，模型及存储保持隔离。"""

from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from scripts.export_chunks import _export_record, _location_label, export
from src.bootstrap import Runtime
from src.chunker import Chunker
from src.config import ExcelSettings, load_settings
from src.context_builder import ContextBuilder
from src.ingestion import list_documents
from src.parsers.excel import ExcelParser
from src.retrieval.hybrid import rrf_fuse
from src.schemas import SearchHit


def parse_excel(tmp_path, workbook, **options):
    """保存真实 XLSX 后使用指定分片配置解析，避免仅模拟单元格对象。"""
    path = tmp_path / "能源统计.XLSX"
    workbook.save(path)
    workbook.close()
    parser = ExcelParser(ExcelSettings(**options), document_factory=SimpleNamespace)
    parsed = parser.parse(path)
    assert not parsed.errors, parsed.errors
    return parsed


def test_excel_repeats_headers_and_keeps_sheet_and_row_locations(tmp_path):
    """每片重复表头，数据行只出现一次；同名数据出现在不同工作表时保留来源。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "电力"
    sheet.append(["项目", "数值"])
    for index in range(5):
        sheet.append([f"项目{index}", index])
    workbook.copy_worksheet(sheet).title = "燃气"
    workbook.create_sheet("空表")
    parsed = parse_excel(tmp_path, workbook, rows_per_chunk=2)
    assert not parsed.texts and not parsed.figures
    assert len(parsed.tables) == 6
    first = parsed.tables[:3]
    assert [(d.metadata["row_start"], d.metadata["row_end"]) for d in first] == [
        (2, 3),
        (4, 5),
        (6, 6),
    ]
    for document in first:
        assert "| 原文行 | 项目 | 数值 |" in document.page_content
        assert document.metadata["sheet_name"] == "电力"
        assert document.metadata["header_row_start"] == 1
        assert document.metadata["column_start"] == "A"
        assert document.metadata["column_end"] == "B"
        assert "page" not in document.metadata
    combined = "\n".join(d.page_content for d in first)
    assert all(combined.count(f"项目{index}") == 1 for index in range(5))


def test_excel_merged_multirow_headers_blank_regions_and_hidden_data(tmp_path):
    """多行表头与合并数据正确展开，空行分隔新区域，隐藏工作表和行也被读取。"""
    workbook = Workbook()
    workbook.active.title = "空白"
    sheet = workbook.create_sheet("隐藏统计")
    sheet.sheet_state = "hidden"
    sheet["B2"] = "能源"
    sheet.merge_cells("B2:C2")
    sheet["B3"], sheet["C3"] = "类别", "数量"
    sheet["B4"], sheet["C4"] = "电力", 10
    sheet.merge_cells("B4:B5")
    sheet["C5"] = 20
    sheet.row_dimensions[5].hidden = True
    sheet["B8"], sheet["B9"], sheet["B10"] = "另一表", "项目", "燃气"
    parsed = parse_excel(tmp_path, workbook, header_rows=2, rows_per_chunk=1)
    assert len(parsed.tables) == 3
    assert "能源 / 类别" in parsed.tables[0].page_content
    assert "能源 / 数量" in parsed.tables[1].page_content
    assert "| 5 | 电力 | 20 |" in parsed.tables[1].page_content
    assert parsed.tables[1].metadata["sheet_state"] == "hidden"
    assert parsed.tables[0].metadata["column_start"] == "B"
    assert parsed.tables[-1].metadata["header_row_start"] == 8


def test_excel_no_header_zero_boolean_dates_and_escaping(tmp_path):
    """无表头模式不丢首行，零、布尔、日期和特殊字符保留明确表示。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([0, False, datetime(2026, 9, 10), "A|B\n第二行"])
    parsed = parse_excel(tmp_path, workbook, header_rows=0)
    content = parsed.tables[0].page_content
    assert "| 原文行 | A列 | B列 | C列 | D列 |" in content
    assert "| 1 | 0 | FALSE | 2026-09-10T00:00:00 | A\\|B<br>第二行 |" in content
    assert "header_row_start" not in parsed.tables[0].metadata


def test_excel_header_only_region_is_not_lost(tmp_path):
    """只有一行文字时仍生成表格，不把被选作表头的内容误判为空。"""
    workbook = Workbook()
    workbook.active["A1"] = "单行说明"
    parsed = parse_excel(tmp_path, workbook)
    assert "单行说明" in parsed.tables[0].page_content
    assert parsed.tables[0].metadata["row_start"] == 1


def test_excel_formulas_keep_cached_zero_and_mark_missing_cache(tmp_path):
    """真实公式缓存为零时保留零，缓存缺失时显示表达式及缺失标记，不执行公式。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["有缓存", "无缓存"])
    sheet.append(["=1-1", "=SUM(1,2)"])
    path = tmp_path / "formula.xlsx"
    workbook.save(path)
    workbook.close()
    # openpyxl 不计算公式；直接在测试文件的 XML 中设置一个合法缓存值模拟 Excel 保存结果。
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    root = ElementTree.fromstring(entries["xl/worksheets/sheet1.xml"])
    namespaces = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    root.find(".//s:c[@r='A2']/s:v", namespaces).text = "0"
    entries["xl/worksheets/sheet1.xml"] = ElementTree.tostring(root)
    with ZipFile(path, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    original = path.read_bytes()
    parsed = ExcelParser(document_factory=SimpleNamespace).parse(path)
    assert not parsed.errors
    content = parsed.tables[0].page_content
    assert "0（缓存结果；公式：=1-1）" in content
    assert "公式：=SUM(1,2)（无缓存结果）" in content
    assert path.read_bytes() == original


@pytest.mark.parametrize("kind", ["missing", "corrupt", "empty", "legacy"])
def test_excel_invalid_files_report_errors(tmp_path, kind):
    """不存在、损坏、无数据和旧 XLS 文件都明确失败。"""
    path = tmp_path / "report.xlsx"
    if kind == "corrupt":
        path.write_bytes(b"not an xlsx")
    elif kind == "empty":
        workbook = Workbook()
        workbook.save(path)
        workbook.close()
    elif kind == "legacy":
        path = path.with_suffix(".xls")
    result = ExcelParser(document_factory=SimpleNamespace).parse(path)
    assert result.errors and not result.tables


def test_excel_oversized_sheet_records_failure_instead_of_truncating(tmp_path):
    """后续工作表超过扫描上限时标记整个文件失败，不把已有的部分结果当作成功。"""
    workbook = Workbook()
    workbook.active["A1"] = "可读取说明"
    workbook.create_sheet("超大范围")["Z100"] = "远端数据"
    path = tmp_path / "large.xlsx"
    workbook.save(path)
    workbook.close()
    result = ExcelParser(
        ExcelSettings(max_sheet_cells=100), document_factory=SimpleNamespace
    ).parse(path)
    assert result.tables and result.errors
    assert "excel.max_sheet_cells" in result.errors[0]


@pytest.mark.parametrize(
    "options",
    [{"header_rows": -1}, {"header_rows": True}, {"rows_per_chunk": 0}, {"max_sheet_cells": 0}],
)
def test_excel_invalid_settings_fail_early(options):
    """错误表头行数和分片上限在解析文件前被拒绝。"""
    with pytest.raises(ValueError):
        ExcelSettings(**options)


def test_excel_settings_are_loaded_and_passed_to_runtime(tmp_path):
    """配置文件的 Excel 参数传到正式解析器，旧配置缺少 Excel 段时使用默认值。"""
    root = Path(__file__).resolve().parents[2]
    content = (root / "config.toml").read_text(encoding="utf-8")
    path = tmp_path / "config.toml"
    path.write_text(content.replace("rows_per_chunk = 50", "rows_per_chunk = 2"), encoding="utf-8")
    runtime = Runtime(load_settings(path))
    workbook = Workbook()
    workbook.active.append(["表头"])
    for number in range(3):
        workbook.active.append([number])
    excel_path = tmp_path / "report.xlsx"
    workbook.save(excel_path)
    workbook.close()
    # 用替身创建文档，实际调用正式注册器中的 ExcelParser。
    runtime.parser._parsers[".xlsx"].document_factory = SimpleNamespace
    assert len(runtime.parser.parse(excel_path).tables) == 2
    before, remainder = content.split("[excel]", 1)
    _, after = remainder.split("[retrieval]", 1)
    path.write_text(before + "[retrieval]" + after, encoding="utf-8")
    assert load_settings(path).excel == ExcelSettings()


def test_excel_citations_export_deduplication_and_chunker(tmp_path):
    """工作表与范围保留到提示词和导出，表格不再语义切分，不同位置的结果不合并。"""
    import csv
    import json

    workbook = Workbook()
    workbook.active.append(["表头"])
    workbook.active.append(["能源数据"])
    parsed = parse_excel(tmp_path, workbook)
    settings = load_settings(Path(__file__).resolve().parents[2] / "config.toml")
    splitter = Mock()
    chunks = Chunker(Mock(), settings.splitting, splitter_factory=splitter).split(parsed)
    splitter.assert_not_called()
    assert chunks == parsed.tables
    original = chunks[0]
    other = SimpleNamespace(
        page_content=original.page_content, metadata={**original.metadata, "sheet_name": "另一表"}
    )
    hits = rrf_fuse([SearchHit(original), SearchHit(other)], [], 10, 60)
    assert len(hits) == 2
    prompt = ContextBuilder("{context}\n{question}").build("数据？", hits).prompt
    assert "工作表: Sheet" in prompt and "第2—2行" in prompt and "表头: 第1—1行" in prompt
    records = [_export_record(original)]
    assert "工作表: Sheet" in _location_label(records[0])
    export(records, "json", tmp_path / "chunks.json")
    assert (
        json.loads((tmp_path / "chunks.json").read_text(encoding="utf-8"))[0]["column_start"] == "A"
    )
    export(records, "csv", tmp_path / "chunks.csv")
    with (tmp_path / "chunks.csv").open(encoding="utf-8-sig", newline="") as stream:
        assert list(csv.DictReader(stream))[0]["row_start"] == "2"


def test_excel_lock_files_and_legacy_format_are_not_discovered(tmp_path):
    """只发现 XLSX 原文件，跳过 Office 锁定文件及旧格式。"""
    for name in ("report.XLSX", "~$report.xlsx", "old.xls"):
        (tmp_path / name).write_bytes(b"placeholder")
    assert [p.name for p in list_documents(tmp_path, {".xlsx"})] == ["report.XLSX"]
